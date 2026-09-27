"""Builds the Google-verified Ormond Beach lead sheet from targets.json + checked.json.
    python3 report.py
Writes ../83apps-ormond-google-verified.xlsx and copies broken-site screenshots to ../ormond-google-broken/.
"""
import json, os, re, shutil, sys
AREA = sys.argv[1] if len(sys.argv) > 1 else 'Ormond Beach'
SLUG = re.sub(r'[^a-z]+', '-', AREA.lower()).strip('-')
OUTDIR = f'../{SLUG}-google-broken' if len(sys.argv) > 1 else '../ormond-google-broken'
XLSX = f'../83apps-{SLUG}-google-verified.xlsx' if len(sys.argv) > 1 else '../83apps-ormond-google-verified.xlsx'
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

SFX = ('-' + sys.argv[2]) if len(sys.argv) > 2 else ''
T = json.load(open(f'targets{SFX}.json'))
C = {c['cid']: c for c in json.load(open(f'checked{SFX}.json'))}
SKIPCAT = re.compile(r'church|government|school|park|library|post office|hospital|city hall|association|non-profit|atm|apartment|condominium|hotel|motel|resort|gas station', re.I)


def maps_link(cid):
    try: return f"https://maps.google.com/?cid={int(cid.split(':')[1], 16)}"
    except Exception: return ''


def cat(r): return ', '.join((r.get('cats') or [])[:2])


def rv(r): return r.get('reviews') or 0


# removed after reviewing screenshots: chains, or a working site where Google links one dead sub-page
DROP = re.compile(r"^(moe's|atlantic animal hospital|critical energy|ppg paint|north american van|dr nawrocki|monica garnache|cathie stull|calvary christian|ifixscreens)", re.I)
broken, check, nosite, social = [], [], [], []
for r in T:
    if SKIPCAT.search(cat(r)): continue
    if r['kind'] == 'own site':
        c = C.get(r['cid'])
        if not c: continue
        if DROP.search(r['name']): continue
        if c['verdict'] == 'BROKEN': broken.append({**r, **c})
        elif c['verdict'] == 'CHECK': check.append({**r, **c})
    elif r['kind'] == 'none' and rv(r) >= 5: nosite.append(r)
    elif r['kind'] in ('social', 'booking/free page') and rv(r) >= 5: social.append(r)
for L in (broken, check, nosite, social): L.sort(key=lambda r: -rv(r))

os.makedirs(OUTDIR, exist_ok=True)
for i, r in enumerate(broken, 1):
    if r.get('shot') and os.path.exists(r['shot']):
        r['shotfile'] = f"{i:02d}-{os.path.basename(r['shot'])}"; shutil.copy(r['shot'], OUTDIR + '/' + r['shotfile'])


def say(r):
    w = r.get('website', '').replace('https://', '').replace('http://', '').rstrip('/')
    y = r.get('why', '').lower()
    if 'www' in w and re.match(r'^www[a-z]', w): return f'"The website link on your Google profile has a typo ({w}), so it never loads." Fix it on the spot, then pitch a proper site.'
    if 'expired' in y or 'exist' in y or 'resolve' in y: return f'"When customers tap Website on your Google listing, they get an error: {w} has expired." Offer to get them back online fast.'
    if 'parked' in y or 'sale' in y: return f'"Your Google listing sends customers to a parked page for {w}." Offer a new site.'
    return f'"Your Google listing\'s website ({w}) shows an error right now." Show the screenshot.'


wb = Workbook()
hdr = Font(bold=True, color='FFFFFF')


def sheet(ws, color, cols, data):
    ws.append([c for c, _, _ in cols])
    for c in ws[1]: c.font = hdr; c.fill = PatternFill('solid', fgColor=color)
    for i, r in enumerate(data, 1): ws.append([f(r, i) for _, f, _ in cols])
    for i, (_, _, w) in enumerate(cols, 1): ws.column_dimensions[get_column_letter(i)].width = w
    for row in ws.iter_rows(min_row=2):
        for c in row: c.alignment = Alignment(wrap_text=True, vertical='top')
    ws.freeze_panes = 'C2'; ws.auto_filter.ref = ws.dimensions


base = [('#', lambda r, i: i, 5), ('Business', lambda r, i: r['name'], 30), ('City', lambda r, i: r.get('city', ''), 16), ('Type', lambda r, i: cat(r), 22),
        ('Google rating', lambda r, i: r.get('rating'), 9), ('Google reviews', lambda r, i: r.get('reviews'), 9)]
tail = [('Address', lambda r, i: (r.get('address') or '').replace(', United States', ''), 34), ('Phone', lambda r, i: r.get('phone'), 15),
        ('Google Maps', lambda r, i: maps_link(r['cid']), 34), ('Contacted?', lambda r, i: '', 11), ('Notes', lambda r, i: '', 30)]
ws = wb.active; ws.title = 'Broken website (verified)'
sheet(ws, '8B1E1E', base + [('Website on Google profile', lambda r, i: r['website'], 34), ('What customers see', lambda r, i: r.get('why'), 40),
                            ('Screenshot', lambda r, i: r.get('shotfile', ''), 26), ('What to say', lambda r, i: say(r), 56)] + tail, broken)
sheet(wb.create_sheet('Check on your phone'), '8A6D1E', base + [('Website on Google profile', lambda r, i: r['website'], 34),
      ('What we saw', lambda r, i: r.get('why'), 40)] + tail, check)
sheet(wb.create_sheet('No website on Google'), '1F1F2A', base + tail, nosite)
sheet(wb.create_sheet('Facebook or booking page only'), '1F4E79', base + [('Link on Google profile', lambda r, i: r.get('website'), 40)] + tail, social)
ws5 = wb.create_sheet('How this was built')
for line in [
    'Source: Google Maps. Each business\'s rating, review count and website link are exactly what its Google profile shows (read 2026-09-27).',
    f'Searched 152 business categories in {AREA}; {len(T)} independent, open businesses (chains, franchises and closed places removed).',
    'Broken website (verified): the website link on the Google profile was tested twice from a real browser. Listed only if the failure is unambiguous:',
    '   domain does not exist / expired, "not found" homepage, parked or for-sale page, expired or suspended hosting, builder "site not connected", or server-down error.',
    'Check on your phone: slow timeouts, repeated server errors, or pages that looked like spam to us. These may be real problems or temporary / network-specific.',
    'No website on Google: the Google profile has no website link at all (5+ reviews).',
    'Tip: tap the Google Maps link in each row, then tap "Website" to see exactly what customers see.',
    'Outreach: personal visits, calls or one-to-one emails only; no automated texts or robocalls (Florida FTSA / TCPA).',
]: ws5.append([line])
ws5.column_dimensions['A'].width = 140
wb.save(XLSX)
print(f'broken {len(broken)}  check {len(check)}  no-website {len(nosite)}  social/booking {len(social)}')
for r in broken[:40]:
    print(f"{r['name'][:34]:34} {r.get('rating')}/{r.get('reviews')}  {r['website'][:40]:40} {r.get('why')}")
