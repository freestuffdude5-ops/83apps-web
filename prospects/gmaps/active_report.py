"""Active-lead report: candidates.json + recency.json -> ../83apps-ACTIVE-leads.xlsx and .csv
Keeps open businesses whose newest visible Google review is within MAX_DAYS.
    python3 active_report.py [max_days]
"""
import json, sys, csv, re
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

MAX_DAYS = int(sys.argv[1]) if len(sys.argv) > 1 else 365
C = json.load(open('candidates.json'))
R = json.load(open('recency.json'))
import os
G = json.load(open('guesses.json')) if os.path.exists('guesses.json') else {}
MALL = re.compile(r'shopping (mall|center|centre|plaza)|plaza|outlet mall|business park|office park|apartment', re.I)
ALREADY = {n.lower() for n in """Smash & Dash|Daytona Blackgold Cycles Inc|FLORIDA FAST TOWING - Local & Long Distance Towing|Precision Elite Athletics|Pro Spray Refinishing LLC|J & M Services|Roy & Company|East Coast Painting|P & P Irrigation and Landscaping|Port Orange MOBiL Complete Auto Repair|Allegiance Tree Care|Port Orange Auto Repair|Higher standards tree care|A1 Septic and Sanitation LLC|GBS Sod & Landscape|Debbie's Mobile Pet Care|Stor-It Boat & RV Center|The Veranda Pampering Salon|Sliding Door Repair Ormond Beach|Osborn Construction & Design|Beachside Dog Grooming|Pirana Grill|Sol Wellness Yoga Studio""".split('|')}


def maps_link(r):
    return f"https://www.google.com/maps/place/?q=place_id:{r['placeId']}" if r.get('placeId') else f"https://maps.google.com/?cid={int(r['cid'].split(':')[1], 16)}"


def age_text(d):
    if d is None: return 'unknown'
    if d < 7: return 'this week'
    if d < 30: return f'~{d // 7} week(s) ago'
    if d < 365: return f'~{d // 30} month(s) ago'
    return f'~{d // 365} year(s) ago'


def say(r):
    lt = r['lead']
    if lt == 'Broken website':
        w = (r.get('website') or '').replace('https://', '').replace('http://', '').rstrip('/')
        return f'"When customers tap Website on your Google listing ({w}), it shows an error." Show them on your phone.'
    if lt == 'No website':
        return f'"You have {r.get("reviews")} Google reviews but nowhere to send people who find you." Offer a simple site with call/quote buttons.'
    return '"Your Google listing only links to a social/booking page." Offer a real site that ranks on Google and links to it.'


active, unknown = [], []
for r in C:
    x = R.get(r['cid'], {})
    if x.get('permClosed') or x.get('tempClosed'): continue
    if MALL.search(', '.join(r.get('cats') or [])) or re.search(r'shopping (center|plaza)|promenade$| plaza$', r['name'], re.I): continue
    r['guess'] = G.get(r['cid'])
    r['newestDays'] = x.get('newestDays'); r['already'] = r['name'].lower() in ALREADY
    if r['newestDays'] is None: unknown.append(r)
    elif r['newestDays'] <= MAX_DAYS: active.append(r)
key = lambda r: -(r.get('reviews') or 0)
cols = [('#', 5), ('Lead type', 16), ('Business', 30), ('Area', 16), ('City', 14), ('Type', 20), ('Google rating', 8), ('Google reviews', 8),
        ('Newest review', 14), ('Website on Google', 30), ('What customers see', 32), ('Address', 32), ('Phone', 15), ('Google Maps', 34),
        ('What to say', 50), ('Check first', 30), ('Already in Client Desk', 10), ('Contacted?', 10), ('Notes', 24)]


def row(i, r):
    return [i, r['lead'], r['name'], r['area'], r.get('city'), ', '.join((r.get('cats') or [])[:2]), r.get('rating'), r.get('reviews'),
            age_text(r.get('newestDays')), r.get('website') or '', r.get('why') or '', (r.get('address') or '').replace(', United States', ''),
            r.get('phone'), maps_link(r), say(r), (f"Probably has a site not linked on Google: {r['guess']} (pitch: fix the Google link, then upgrade)" if r.get('guess') else ''), 'yes' if r['already'] else '', '', '']


wb = Workbook(); first = True
for title, lead in [('Broken website (active)', 'Broken website'), ('No website (active)', 'No website'), ('Facebook-booking only (active)', 'Facebook/booking page only')]:
    ws = wb.active if first else wb.create_sheet(); first = False; ws.title = title
    data = sorted([r for r in active if r['lead'] == lead], key=key)
    ws.append([c for c, _ in cols])
    for c in ws[1]: c.font = Font(bold=True, color='FFFFFF'); c.fill = PatternFill('solid', fgColor={'Broken website': '8B1E1E', 'No website': '1F1F2A'}.get(lead, '1F4E79'))
    for i, r in enumerate(data, 1): ws.append(row(i, r))
    for i, (_, w) in enumerate(cols, 1): ws.column_dimensions[get_column_letter(i)].width = w
    for rr in ws.iter_rows(min_row=2):
        for c in rr: c.alignment = Alignment(wrap_text=True, vertical='top')
    ws.freeze_panes = 'D2'; ws.auto_filter.ref = ws.dimensions
ws = wb.create_sheet('Activity unknown'); ws.append([c for c, _ in cols])
for i, r in enumerate(sorted(unknown, key=key), 1): ws.append(row(i, r))
ws2 = wb.create_sheet('How this was built')
for line in [f'Active = open on Google and the newest review we could see is within {MAX_DAYS} days.',
             '"Newest review" is the most recent of the reviews Google showed on the profile (Google does not show us all reviews).',
             'Broken website leads were tested twice from a real browser and reviewed by screenshot; No website = Google profile has no website link.',
             'Tap the Google Maps link to confirm before visiting. Outreach: personal visits, calls or one-to-one emails only.']:
    ws2.append([line])
wb.save('../83apps-ACTIVE-leads.xlsx')
with open('../83apps-ACTIVE-leads.csv', 'w', newline='') as fh:
    w = csv.writer(fh); w.writerow([c for c, _ in cols])
    for i, r in enumerate(sorted(active, key=lambda r: (r['lead'] != 'Broken website', -(r.get('reviews') or 0))), 1): w.writerow(row(i, r))
from collections import Counter
print('active', len(active), Counter(r['lead'] for r in active), '| unknown', len(unknown), '| already', sum(r['already'] for r in active))
for lead in ('Broken website', 'No website'):
    print('==', lead)
    for r in sorted([r for r in active if r['lead'] == lead and not r['already']], key=key)[:12]:
        print(f"  {r['name'][:34]:34} {r['area'][:12]:12} {r.get('rating')}/{r.get('reviews')}  newest {age_text(r['newestDays'])}")
