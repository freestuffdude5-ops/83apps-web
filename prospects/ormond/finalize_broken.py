"""Builds the "broken website" shortlist from broken_verified.json + ranked.json.
    python3 finalize_broken.py
Adds a first sheet to ../83apps-ormond-prospects.xlsx and copies evidence screenshots to ../ormond-broken-sites/.
"""
import json, re, os, shutil
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

ver = {r['name']: r for r in json.load(open('broken_verified.json'))}
ranked = {r['name']: r for r in json.load(open('ranked.json'))}

# what a customer sees, confirmed by eye from the screenshots (overrides the automatic label)
SEEN = {
    'Beauty Nails & Spa': 'Link on their listing opens an unrelated blank "HTTP Server Test Page"',
    'Arborist Pros Tree Service LLC': 'Domain expired: registrar page says "has expired, Renew now"',
    'Bliss Massage Spa': 'Domain suspended by the host ("Domain Temporarily Disabled")',
    "Simone's Salon": 'Site down: "The request could not be satisfied" server error',
    'Coastal Appliance Service (Coastal Appliance Repair Services Inc.)': 'Old domain now an ad sending callers to a DIFFERENT appliance-repair number',
    'Artistic Soul Tattoo': 'Wix error: "This domain isn\'t connected to a site"',
    "Austin's Mobile Marine": 'Wix error: "This domain isn\'t connected to a site"',
    'Sharps Discount Liquors': 'Domain hijacked: shows Chinese gambling spam',
    'Veranda Pampering Salon': 'Cloudflare error 521 "Web server is down"',
    'Utopian Photos': 'Squarespace "Website Expired" page',
    "Owen Heating & Cooling Inc.": 'Site will not load (security/certificate error)',
    'Molto Bella Boutique': 'Domain hijacked: redirects to an Indonesian gambling site',
    'Thai Wood House (Thai Woodhouse)': 'Domain hijacked: redirects to a gambling site',
    "Lucy's Gift Boutique": 'Domain hijacked: shows a slot-gambling site',
    'The Flower Market (Ormond Beach Flower Market)': 'Domain hijacked: redirects to a gambling site',
    'The Pocket Jeweler': 'Domain hijacked: redirects to an ADULT website (no screenshot saved)',
    'Beach Village Gift Shop (Beach Village Gift Emporium)': 'Squarespace "Coming Soon, under construction" page',
    'Fugu Sushi': 'Cloudflare error 522 "Connection timed out"',
    'Pirana Grille': 'Cloudflare "DNS resolution error": site is gone',
}
CONFIRM = {'Southern Auto Source', 'Kimble Electric Co.', 'Tire City (Tire City of Volusia County)', 'Browns Electric of Central Florida',
           'Carey Plumbing Inc.', 'Lars Air, LLC', 'Air One Heating & Cooling', 'Buckels Sprinklers',
           "Don Pepper's Mexican Grill & Cantina", 'Dog Hut', 'Joyologie Boutique'}
DROP = {"Diane's Pet Grooming": 'site works (one-page site with contact form)', 'Thai Erawan': 'thaierawanrestaurant.com works'}


def is_confirm(name):
    return any(name.startswith(c.split(' (')[0]) for c in CONFIRM)


def listed_where(r):
    t = ' '.join(str(r.get(k) or '') for k in ('websiteStatus', 'websiteEvidence', 'notes'))
    where = [w for w, rx in [('Google profile', r'google (profile|listing|maps)|on google|listing.s website|gbp'), ('Yelp', r'yelp'),
                             ('Nextdoor', r'nextdoor'), ('YellowPages', r'yellowpages|\byp\b'), ('Birdeye', r'birdeye'), ('Facebook', r'facebook')]
             if re.search(rx, t, re.I)]
    return ', '.join(where) or 'directory listings'


rows = []
for name, v in ver.items():
    if name in DROP: continue
    r = ranked.get(name, {})
    seen = SEEN.get(name)
    if not seen:
        base = v['verdict']
        seen = {'Domain does not exist (expired or never renewed)': 'Domain expired: browsers show "server can\'t be found"'}.get(base, base)
        if is_confirm(name):
            seen = (r.get('websiteStatus') or base).split(':', 1)[-1].strip()[:90] + ' (confirm on your phone)'
    rows.append(dict(name=name, r=r, v=v, seen=seen, confirm=is_confirm(name)))


def n(x):
    c = x['r'].get('googleReviewCount')
    return c if isinstance(c, (int, float)) else 0


rows.sort(key=lambda x: (x['confirm'], -n(x)))
os.makedirs('../ormond-broken-sites', exist_ok=True)
for i, x in enumerate(rows, 1):
    s = x['v'].get('shot')
    if s and os.path.exists(s) and not x['confirm'] and ('LOADS' not in x['v']['verdict'] or x['name'] in SEEN):
        dst = f"../ormond-broken-sites/{i:02d}-{os.path.basename(s)}"
        shutil.copy(s, dst); x['shotfile'] = os.path.basename(dst)


def pitch(x):
    s = x['seen'].lower(); url = x['v']['url'].replace('http://', '').replace('https://', '').rstrip('/')
    if 'different' in s: return f'"Your old site {url} is now an ad sending your callers to another repair company." Offer a new site this week.'
    if 'hijack' in s or 'unrelated' in s: return f'"The website link people click for you goes to a page that has nothing to do with you." Show the screenshot.'
    if 'expired' in s or "can't be found" in s: return f'"Your website {url} expired, so anyone who clicks it gets an error." Offer to get it back up fast.'
    return f'"Your website link ({url}) shows an error page right now." Show the screenshot and offer a fast fix.'


wb = load_workbook('../83apps-ormond-prospects.xlsx')
if 'Broken websites' in wb.sheetnames: del wb['Broken websites']
ws = wb.create_sheet('Broken websites', 0)
cols = [('#', 5), ('Business', 30), ('Type', 20), ('Google rating', 9), ('Google reviews', 9), ('Other reviews', 20), ('Broken link', 30), ('What customers see', 46),
        ('Link appears on', 22), ('Screenshot', 30), ('Address', 30), ('Phone', 15), ('What to say', 56), ('Contacted?', 11), ('Notes', 30)]
ws.append([c for c, _ in cols])
hdr = Font(bold=True, color='FFFFFF'); fill = PatternFill('solid', fgColor='8B1E1E')
for c in ws[1]: c.font = hdr; c.fill = fill
for i, x in enumerate(rows, 1):
    r = x['r']
    ws.append([i, x['name'], r.get('category'), r.get('googleRating'), r.get('googleReviewCount'), r.get('otherReviews'), x['v']['url'], x['seen'], listed_where(r),
               x.get('shotfile', ''), r.get('address'), r.get('phone'), pitch(x), '', r.get('notes')])
for i, (_, w) in enumerate(cols, 1): ws.column_dimensions[get_column_letter(i)].width = w
for row in ws.iter_rows(min_row=2):
    for c in row: c.alignment = Alignment(wrap_text=True, vertical='top')
ws.freeze_panes = 'C2'
ws.auto_filter.ref = ws.dimensions
wb.save('../83apps-ormond-prospects.xlsx')
json.dump([{k: v for k, v in x.items() if k not in ('r', 'v')} | {'url': x['v']['url']} for x in rows], open('broken_final.json', 'w'), indent=1)
print(len(rows), 'broken-site prospects;', sum(not x['confirm'] for x in rows), 'confirmed from here')
for i, x in enumerate(rows[:25], 1):
    print(f"{i:>2}. {x['name'][:34]:34} G {x['r'].get('googleRating')}/{x['r'].get('googleReviewCount')}  {x['seen'][:60]}")
