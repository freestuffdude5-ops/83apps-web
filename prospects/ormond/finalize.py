"""Merges the Ormond Beach research files, dedupes, re-checks websites, ranks, and writes the spreadsheet.
    python3 finalize.py
Inputs: *.json research files in this folder (home-services, auto-marine, beauty-pets, trades-services,
food-retail, osm-verified). Output: ../83apps-ormond-prospects.xlsx and ranked.json
"""
import json, re, math, glob, subprocess
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

SRC = {'home-services.json': ('Home services', 1.0), 'auto-marine.json': ('Auto & marine', 0.95),
       'beauty-pets.json': ('Beauty & pets', 0.95), 'trades-services.json': ('Trades & services', 0.95),
       'food-retail.json': ('Food & retail', 0.8), 'osm-verified.json': (None, 0.9)}


def norm(s):
    return re.sub(r'[^a-z0-9]', '', (s or '').lower().replace('&', 'and').replace('llc', '').replace('inc', ''))


def digits(p):
    d = re.sub(r'\D', '', p or '')
    return d[-10:] if len(d) >= 10 else None


def guess_group(cat):
    c = (cat or '').lower()
    if re.search(r'restaurant|cafe|coffee|food|bar|pub|bakery|ice cream|deli|pizza|diner|grill|shop|store|boutique|gift|thrift|florist|retail|clothes|second', c):
        return 'Food & retail', 0.8
    if re.search(r'car|auto|tire|tyre|boat|marine|motorcycle|golf', c): return 'Auto & marine', 0.95
    if re.search(r'hair|barber|nail|beauty|salon|spa|massage|lash|tattoo|pet|groom|dog', c): return 'Beauty & pets', 0.95
    if re.search(r'lawn|landscap|clean|pool|pest|handyman|paint|roof|fence', c): return 'Home services', 1.0
    return 'Trades & services', 0.9


rows = []
for f, (group, w) in SRC.items():
    try:
        data = json.load(open(f))
    except FileNotFoundError:
        print('missing', f); continue
    for r in data:
        r['_group'], r['_w'] = (group, w) if group else guess_group(r.get('category'))
        r['_file'] = f
        rows.append(r)

# dedupe on phone or normalized name; keep the richer record
seen = {}
for r in rows:
    keys = [k for k in (digits(r.get('phone')), norm(r.get('name'))) if k]
    hit = next((seen[k] for k in keys if k in seen), None)
    if hit is None:
        for k in keys: seen[k] = r
        continue
    richer = r if len(json.dumps(r)) > len(json.dumps(hit)) else hit
    other = hit if richer is r else r
    richer.setdefault('_alsoIn', []).append(other['_file'])
    for k in keys + [digits(hit.get('phone')), norm(hit.get('name'))]:
        if k: seen[k] = richer
uniq = list({id(v): v for v in seen.values()}.values())

# re-check any website URL an agent found (dead / free subdomain / placeholder), so the sheet reflects today
URL = re.compile(r'https?://[^\s,;)"]+')


def check(url):
    try:
        out = subprocess.run(['curl', '-sSL', '-m', '20', '-o', '/dev/null', '-w', '%{http_code} %{url_effective}', url],
                             capture_output=True, text=True, timeout=30).stdout.strip()
    except Exception as e:
        out = 'ERR'
    return out


for r in uniq:
    st = r.get('websiteStatus') or ''
    m = URL.search(st)
    if m and not st.startswith(('third-party', 'booking', 'marketplace')):
        r['_recheck'] = f'{m.group(0)} -> {check(m.group(0))}'


def reviews(r):
    n = r.get('googleReviewCount')
    if isinstance(n, (int, float)) and n > 0: return int(n), 'Google'
    m = re.search(r'(\d[\d,]*)\s*(?:reviews|ratings|recommend|\))', str(r.get('otherReviews') or ''), re.I)
    if m: return int(m.group(1).replace(',', '')), 'other'
    return 0, None


def score(r):
    n, src = reviews(r)
    s = min(math.log1p(n) / math.log1p(150), 1) * 50
    rt = r.get('googleRating')
    if isinstance(rt, (int, float)):
        s += 15 if rt >= 4.5 else 10 if rt >= 4.0 else 3
    st = (r.get('websiteStatus') or '').lower()
    s += 30 if st.startswith('dead') else 25 if st.startswith('none') else 20
    s += 10 if r.get('confidence') == 'high' else 3
    if src != 'Google': s -= 5
    return round(s * r['_w'], 1)


# removed after re-check: now has its own working domain
EXCLUDE = {'justincredibledetailing': 'wixsite now redirects to its own domain justin-credible-detailing.com'}
uniq = [r for r in uniq if norm(r.get('name')) not in EXCLUDE]
for r in uniq:
    r['_score'] = score(r); r['_reviews'] = reviews(r)
ranked = sorted(uniq, key=lambda r: -r['_score'])
json.dump(ranked, open('ranked.json', 'w'), indent=1, default=str)


def pitch(r):
    st = (r.get('websiteStatus') or '').lower(); g = r['_group']
    if st.startswith('dead'): return 'Their website link is broken: people who find them on Google hit an error. Offer a fast replacement.'
    if 'free subdomain' in st or 'placeholder' in st: return 'They started a site but it is a free placeholder. Offer a real site on their own domain.'
    if 'booking' in st or 'third-party' in st or 'marketplace' in st:
        return 'They rely on a booking/ordering page. Offer a real site that ranks on Google and links to it.'
    if g in ('Home services', 'Trades & services', 'Auto & marine'):
        return 'Great reviews but nowhere to send Google searchers: site with tap-to-call + quote form, then review requests.'
    if g == 'Beauty & pets': return 'Reviews but no site: simple site with online booking and reminders.'
    return 'Reviews but no site: menu/hours/photos site that shows up on Google, with ordering links.'


wb = Workbook(); ws = wb.active; ws.title = 'Ormond Beach - no website'
hdr = Font(bold=True, color='FFFFFF'); fill = PatternFill('solid', fgColor='1F1F2A')
cols = [('#', 5), ('Business', 30), ('Type', 20), ('Group', 15), ('Google rating', 9), ('Google reviews', 9), ('Other reviews', 22),
        ('Website status', 34), ('Address', 30), ('Phone', 15), ('Facebook / links', 34), ('Why it qualifies', 46), ('Pitch angle', 46),
        ('Confidence', 10), ('Contacted?', 11), ('Notes', 30)]
ws.append([c for c, _ in cols])
for c in ws[1]: c.font = hdr; c.fill = fill
for i, r in enumerate(ranked, 1):
    why = '; '.join(x for x in [r.get('websiteEvidence'), r.get('activeSignals'), r.get('_recheck') and 'Re-checked: ' + r['_recheck']] if x)
    ws.append([i, r.get('name'), r.get('category'), r['_group'], r.get('googleRating'), r.get('googleReviewCount'), r.get('otherReviews'),
               r.get('websiteStatus'), r.get('address'), r.get('phone'), ' '.join(x for x in [r.get('facebook'), r.get('otherLinks') if isinstance(r.get('otherLinks'), str) else ' '.join(r.get('otherLinks') or [])] if x),
               why, pitch(r), r.get('confidence'), '', r.get('notes')])
for i, (_, wdt) in enumerate(cols, 1): ws.column_dimensions[get_column_letter(i)].width = wdt
for row in ws.iter_rows(min_row=2):
    for c in row: c.alignment = Alignment(wrap_text=True, vertical='top')
ws.freeze_panes = 'C2'
ws.auto_filter.ref = ws.dimensions
ws2 = wb.create_sheet('How this was built')
for line in [
    'Ormond Beach / Ormond-by-the-Sea small businesses that have customer reviews but no real website of their own.',
    '"No real website" = none (Facebook/Instagram/Google only), a dead link, a free builder subdomain or placeholder, or only a booking/ordering page.',
    'Sources: web search, Yelp, Nextdoor, Facebook, BBB, directories that republish Google ratings, and OpenStreetMap (checked by hand).',
    'Google rating/review counts are only filled in where they were seen in a source; blank means not visible, not zero.',
    'Rank = review strength (count + rating) + how clearly they lack a site + confidence, weighted by how well the business type fits.',
    'Before outreach: open their Google listing to confirm they are open and still have no website linked.',
    'Outreach: personal calls, visits, or one-to-one emails only. No automated texts/robocalls (Florida FTSA / federal TCPA).',
]: ws2.append([line])
ws2.column_dimensions['A'].width = 140
wb.save('../83apps-ormond-prospects.xlsx')
print('merged', len(rows), 'unique', len(uniq))
for r in ranked[:25]:
    print(f"{r['_score']:>5} {r.get('name','')[:36]:36} {r['_group'][:14]:14} G:{r.get('googleRating')}/{r.get('googleReviewCount')} {str(r.get('websiteStatus'))[:40]}")
